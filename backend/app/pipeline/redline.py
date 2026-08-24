"""Degisiklik izlemeli (tracked changes) redline belgesi uretimi.

Hukukcunun bekledigi format budur: yanina yazilmis bir oneri degil, Word'de
KABUL ET / REDDET yapabilecegi gercek bir degisiklik.

python-docx tracked changes desteklemez; w:ins / w:del ogelerini dogrudan
belgenin XML'ine yaziyoruz.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

AUTHOR = "Sözleşme Feneri"
_counter = {"id": 1000}


def _next_id() -> str:
    _counter["id"] += 1
    return str(_counter["id"])


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _run(text: str, tag: str = "w:t") -> OxmlElement:
    r = OxmlElement("w:r")
    t = OxmlElement(tag)
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add_deletion(paragraph, text: str) -> None:
    """Silinecek metni degisiklik izlemeli olarak isaretler."""
    if not text:
        return
    d = OxmlElement("w:del")
    d.set(qn("w:id"), _next_id())
    d.set(qn("w:author"), AUTHOR)
    d.set(qn("w:date"), _stamp())
    d.append(_run(text, "w:delText"))
    paragraph._p.append(d)


def add_insertion(paragraph, text: str) -> None:
    """Eklenecek metni degisiklik izlemeli olarak isaretler."""
    if not text:
        return
    i = OxmlElement("w:ins")
    i.set(qn("w:id"), _next_id())
    i.set(qn("w:author"), AUTHOR)
    i.set(qn("w:date"), _stamp())
    i.append(_run(text, "w:t"))
    paragraph._p.append(i)


def add_comment_note(doc, text: str) -> None:
    """Gerekce notu - degisiklik degil, aciklama."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(14)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x6E, 0x77, 0x89)


def build_redline(payload: dict, out_path: Path) -> Path:
    """Onerilen madde metinlerini degisiklik izlemeli tek belgede toplar."""
    c = payload["contract"]
    findings = [
        f for f in payload["findings"]
        if f.get("proposed_text") and f["finding_type"] != "INFO"
    ]

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    h = doc.add_paragraph()
    hr = h.add_run("SÖZLEŞME DEĞİŞİKLİK ÖNERİLERİ")
    hr.bold = True
    hr.font.size = Pt(18)

    s = doc.add_paragraph()
    sr = s.add_run(
        f"{c.get('filename')} · {c.get('counterparty') or 'karşı taraf belirtilmemiş'} · "
        f"{payload.get('generated_at')}"
    )
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x6E, 0x77, 0x89)

    intro = doc.add_paragraph()
    ir = intro.add_run(
        "Bu belge değişiklik izleme (track changes) biçimindedir. Word'de "
        "Gözden Geçir sekmesinden her değişikliği ayrı ayrı kabul edebilir veya "
        "reddedebilirsiniz. Üstü çizili metin mevcut sözleşmedeki hâli, altı çizili "
        "metin Banka'nın önerdiği hâli gösterir."
    )
    ir.font.size = Pt(9)
    doc.add_paragraph()

    if not findings:
        doc.add_paragraph("Değişiklik önerisi üretilmedi.")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return out_path

    order = {"KRITIK": 0, "YUKSEK": 1, "ORTA": 2, "DUSUK": 3, "BILGI": 4}
    findings.sort(key=lambda f: (order.get(f["severity"], 9), f.get("clause_number", "")))

    for n, f in enumerate(findings, 1):
        loc = f"Madde {f['clause_number']}" if f.get("clause_number") else "YENİ MADDE (sözleşmede yok)"
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(14)
        head.paragraph_format.space_after = Pt(2)
        hr2 = head.add_run(f"{n}. {loc} — {f['title']}")
        hr2.bold = True
        hr2.font.size = Pt(11)

        tag = doc.add_paragraph()
        tag.paragraph_format.space_after = Pt(4)
        tr = tag.add_run(f"[{f['severity']}]  {f['code']}")
        tr.font.size = Pt(8)
        tr.font.color.rgb = RGBColor(0x6E, 0x77, 0x89)

        body = doc.add_paragraph()
        body.paragraph_format.space_after = Pt(4)
        # Baglam alintisi SILINECEK metin degildir; onu silme olarak isaretlemek
        # hukukcuya yanlis bir degisiklik onerir.
        if f.get("quote") and f.get("quote_is_evidence", True):
            # Mevcut metin siliniyor, onerilen metin ekleniyor.
            add_deletion(body, f["quote"].strip())
            add_insertion(body, " " + f["proposed_text"].strip())
        else:
            # Eksik madde: yalnizca ekleme.
            add_insertion(body, f["proposed_text"].strip())

        note = f.get("rationale") or ""
        if f.get("legal_basis"):
            note += "  Dayanak: " + "; ".join(f["legal_basis"])
        if f.get("negotiation_note"):
            note += "  Müzakere notu: " + f["negotiation_note"]
        add_comment_note(doc, note.strip())

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
