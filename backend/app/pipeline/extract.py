"""Asama 2: Metin cikarma (PDF / DOCX / TXT)."""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

from . import ocr as ocr_mod

log = logging.getLogger(__name__)

HEADING_MARK = "[[H]]"


class ExtractionError(RuntimeError):
    pass


@dataclass
class ExtractSonuc:
    text: str
    pages: list[dict] = field(default_factory=list)
    ocr_used: bool = False
    note: str = ""
    ocr_confidence: float | None = None


def extract_ex(path: str | Path) -> ExtractSonuc:
    """Metin cikarimi; PDF'te metin katmani yetersizse OCR'a duser."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _pdf(p)
    if suffix == ".docx":
        t, pages = _docx(p)
        return ExtractSonuc(text=t, pages=pages, note="DOCX metin katmanı")
    if suffix in (".txt", ".md"):
        t = p.read_text(encoding="utf-8", errors="replace")
        return ExtractSonuc(text=t, pages=[{"page": 1, "start": 0, "end": len(t)}],
                            note="düz metin")
    raise ExtractionError(f"Desteklenmeyen dosya türü: {suffix}")


def extract(path: str | Path) -> tuple[str, list[dict]]:
    """Geriye donuk sade arayuz: (ham_metin, sayfa_haritasi)."""
    r = extract_ex(path)
    return r.text, r.pages


def _pdf(p: Path) -> ExtractSonuc:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    parts: list[str] = []
    pages: list[dict] = []
    cursor = 0
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception:  # bozuk sayfa tum belgeyi dusurmesin
            txt = ""
        txt = txt.replace("\r\n", "\n")
        parts.append(txt)
        pages.append({"page": i, "start": cursor, "end": cursor + len(txt)})
        cursor += len(txt) + 1
    text = "\n".join(parts)

    if not ocr_mod.metin_katmani_yetersiz(text, len(pages)):
        return ExtractSonuc(text=text, pages=pages, note=f"PDF metin katmanı ({len(pages)} sayfa)")

    # --- metin katmani yok: taranmis belge ---
    hazir, neden = ocr_mod.kullanilabilir()
    if not hazir:
        raise ExtractionError(
            "PDF içinde okunabilir metin katmanı yok (taranmış belge) ve OCR kullanılamıyor: "
            f"{neden}. Çözüm: sunucuya tesseract kurun "
            "(Debian/Ubuntu: apt-get install tesseract-ocr tesseract-ocr-tur) "
            "veya belgeyi DOCX ya da metin katmanlı PDF olarak yükleyin."
        )

    log.info("PDF'te metin katmanı yok, OCR devreye giriyor: %s", p.name)
    try:
        r = ocr_mod.pdf_ocr(p)
    except ocr_mod.OCRUnavailable as exc:
        raise ExtractionError(str(exc)) from exc
    return ExtractSonuc(text=r.text, pages=r.pages, ocr_used=True,
                        note=r.note, ocr_confidence=r.mean_confidence)


def _docx(p: Path) -> tuple[str, list[dict]]:
    import docx  # python-docx

    d = docx.Document(str(p))
    lines: list[str] = []
    for para in d.paragraphs:
        t = (para.text or "").strip()
        style = (para.style.name or "") if para.style else ""
        # Baslik stilleri madde tespitine yardimci olsun diye isaretlenir.
        if t and style.lower().startswith(("heading", "baslik")):
            lines.append(HEADING_MARK + t)
        else:
            lines.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    if len(text.strip()) < 100:
        raise ExtractionError("DOCX içinde anlamlı metin bulunamadı.")
    return text, [{"page": 1, "start": 0, "end": len(text)}]
